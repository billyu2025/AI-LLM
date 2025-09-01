import csv
import os
from pathlib import Path

from docx import Document
from docx.shared import Pt


def ensure_dir(path: str) -> None:
    Path(path).mkdir(parents=True, exist_ok=True)


def sanitize_for_xml(text: str) -> str:
    if text is None:
        return ""
    # Remove control chars not allowed in XML (except common whitespace) and NULL bytes
    return "".join(
        ch for ch in str(text)
        if (
            ch == "\t" or ch == "\n" or ch == "\r" or ord(ch) >= 32
        )
    )


def add_markdown_as_paragraphs(document: Document, md_text: str) -> None:
    lines = md_text.splitlines()
    for raw_line in lines:
        line = raw_line.rstrip("\n")
        if not line.strip():
            document.add_paragraph("")
            continue

        # Headings based on simple heuristics
        if line.lower().startswith("slide ") or line.endswith("/ 核心概要") or "—" in line and line.strip().lower().startswith("slide"):
            document.add_heading(line.strip(), level=1)
            continue

        # Section separator
        if line.strip() in {"—", "---", "——"}:
            document.add_paragraph("—" * 20)
            continue

        # Bullets
        if line.lstrip().startswith("- "):
            p = document.add_paragraph(line.lstrip()[2:].strip())
            p.style = "List Bullet"
            continue

        # Default paragraph
        document.add_paragraph(line)


def export_markdown_to_docx(md_path: str, docx_path: str) -> None:
    with open(md_path, "r", encoding="utf-8") as f:
        md_text = f.read()
    doc = Document()
    style = doc.styles["Normal"].font
    style.name = "Calibri"
    style.size = Pt(11)
    add_markdown_as_paragraphs(doc, md_text)
    doc.save(docx_path)


def export_csv_to_docx(csv_path: str, docx_path: str, max_rows: int | None = None) -> None:
    doc = Document()
    doc.add_heading("HK Retail/F&B/Hospitality Prospects", level=1)

    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        fieldnames = reader.fieldnames or []
        rows = list(reader)

    if max_rows is not None:
        rows = rows[:max_rows]

    table = doc.add_table(rows=1 + len(rows), cols=len(fieldnames))
    hdr_cells = table.rows[0].cells
    for i, name in enumerate(fieldnames):
        hdr_cells[i].text = name

    for r_idx, row in enumerate(rows, start=1):
        for c_idx, name in enumerate(fieldnames):
            value = sanitize_for_xml(row.get(name, ""))
            table.rows[r_idx].cells[c_idx].text = value

    doc.save(docx_path)


def export_combined_docx(parts: list[tuple[str, str]], docx_path: str) -> None:
    doc = Document()
    style = doc.styles["Normal"].font
    style.name = "Calibri"
    style.size = Pt(11)

    for title, md_path in parts:
        doc.add_heading(title, level=1)
        with open(md_path, "r", encoding="utf-8") as f:
            add_markdown_as_paragraphs(doc, f.read())
        doc.add_page_break()

    doc.save(docx_path)


def main() -> None:
    root = "/workspace"
    exports_dir = os.path.join(root, "sales", "exports")
    ensure_dir(exports_dir)

    deck_md = os.path.join(root, "sales", "customer-service-copilot", "Customer-Service-Copilot-Deck-zh-en.md")
    demo_md = os.path.join(root, "sales", "customer-service-copilot", "Demo-Script-zh-en.md")
    csv_path = os.path.join(root, "sales", "prospects", "hk-retail-fnb-hospitality-100.csv")

    deck_docx = os.path.join(exports_dir, "Customer-Service-Copilot-Deck-zh-en.docx")
    demo_docx = os.path.join(exports_dir, "Demo-Script-zh-en.docx")
    combined_docx = os.path.join(exports_dir, "CSC-Deck-and-Demo-zh-en.docx")
    list_docx = os.path.join(exports_dir, "hk-retail-fnb-hospitality-list.docx")

    export_markdown_to_docx(deck_md, deck_docx)
    export_markdown_to_docx(demo_md, demo_docx)
    export_combined_docx([
        ("Customer Service Copilot — Deck", deck_md),
        ("Customer Service Copilot — Demo Script", demo_md),
    ], combined_docx)

    # Export current version of CSV list (all rows present)
    if os.path.exists(csv_path):
        export_csv_to_docx(csv_path, list_docx)

    print("Exported:")
    print(deck_docx)
    print(demo_docx)
    print(combined_docx)
    if os.path.exists(csv_path):
        print(list_docx)


if __name__ == "__main__":
    main()

